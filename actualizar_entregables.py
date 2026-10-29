"""Genera los entregables del avance (ficha técnica, Gantt y presentación) con las cifras de esta ejecución.

Uso (desde TA_E10, después de `python run_all.py`):  python actualizar_entregables.py

Parte de las versiones originales en entregables/plantillas/ (mismo diseño) y reemplaza textos y figuras.
Cada reemplazo debe encontrar su texto de origen; si no lo encuentra, el script se detiene.
Las cifras se leen de datos/4_resultados/cifras_*.json, que escribe cada notebook.
"""
import copy, json, shutil
from pathlib import Path
import openpyxl
from pptx import Presentation
from pptx.util import Emu

RAIZ = Path(__file__).resolve().parent
PLANT, ENT, FIG = RAIZ / "entregables" / "plantillas", RAIZ / "entregables", RAIZ / "figuras"
C = {}
for f in sorted((RAIZ / "datos" / "4_resultados").glob("cifras_*.json")):
    C[f.stem.replace("cifras_", "")] = json.loads(f.read_text())


def n(x, d=1):
    """Número con coma decimal (estilo del informe)."""
    return f"{x:.{d}f}".replace(".", ",")


s1, s2a, s2b, s3, s4, s5, s6, a4, a6 = (C[k] for k in ["step1", "step2a", "step2b", "step3", "step4", "step5", "step6", "anexo4", "anexo6"])
RANGO_LOGIT = f"{n(s6['auc_semillas_min'], 2)}–{n(s6['auc_semillas_max'], 2)}"
RANGO_SVC = f"{n(s6['auc_svc_semillas_min'], 2)}–{n(s6['auc_svc_semillas_max'], 2)}"
EVAL = f"{s3['evaluaciones']} evaluaciones, {s3['ninos_eval']} NNA"


# ---------------------------------------------------------------- PPTX
def reemplazar_parrafos(frame, viejo, nuevo):
    hechos = 0
    for p in frame.paragraphs:
        runs = p.runs
        texto = "".join(r.text for r in runs)
        if viejo not in texto:
            continue
        hit = next((r for r in runs if viejo in r.text), None)
        if hit is not None:
            hit.text = hit.text.replace(viejo, nuevo)
        else:  # el texto está repartido entre varios runs: se conserva el formato del primero
            runs[0].text = texto.replace(viejo, nuevo)
            for r in runs[1:]:
                r.text = ""
        hechos += 1
    return hechos


def reemplazar(slide, viejo, nuevo):
    hechos = 0
    for sh in slide.shapes:
        if sh.has_text_frame:
            hechos += reemplazar_parrafos(sh.text_frame, viejo, nuevo)
        elif sh.has_table:
            for fila in sh.table.rows:
                for celda in fila.cells:
                    hechos += reemplazar_parrafos(celda.text_frame, viejo, nuevo)
    if hechos == 0:
        raise SystemExit(f"No se encontró en la diapositiva: {viejo!r}")


def cambiar_imagen(slide, nombre_shape, archivo):
    """Reemplaza la imagen y la encaja, sin deformarla, dentro de la caja de la imagen original."""
    pic = next(sh for sh in slide.shapes if sh.name == nombre_shape)
    from PIL import Image
    w, h = Image.open(archivo).size
    rid = pic._element.blipFill.blip.rEmbed
    pic.part.related_part(rid)._blob = Path(archivo).read_bytes()
    escala = min(pic.width / w, pic.height / h)
    pic.width, pic.height = Emu(int(w * escala)), Emu(int(h * escala))


def agregar_vineta(slide, nombre_shape, texto):
    sh = next(s for s in slide.shapes if s.name == nombre_shape)
    ultimo = sh.text_frame.paragraphs[-1]._p
    nuevo = copy.deepcopy(ultimo); ultimo.addnext(nuevo)
    p = sh.text_frame.paragraphs[-1]
    p.runs[0].text = texto
    for r in p.runs[1:]: r.text = ""


def pptx():
    prs = Presentation(PLANT / "Avance_TA_E10_avance_original.pptx")
    S = prs.slides
    R = lambda i, *pares: [reemplazar(S[i - 1], a, b) for a, b in pares]
    R(2, ("Ficha y medición OMS concuerdan en 68,6 %. Con las medidas se estima la alteración estricta con AUC ≈ 0,82; con la ficha sola, no (AUC ≈ 0,5).",
          f"Ficha y medición OMS concuerdan en {n(s5['concordancia'])} %. Con las medidas se estima la alteración estricta con AUC {RANGO_LOGIT} (según la partición); con la ficha sola, no (AUC ≈ 0,5)."))
    R(3, ("Resultado: AUC 0,41 a 0,51, sin diferencia con el azar (p 0,80 a 0,99).",
          f"Resultado: AUC {n(a6['auc_min_modelos'], 2)} a {n(a6['auc_max_modelos'], 2)}, sin diferencia con el azar (p {n(a6['p_min_permutaciones'], 2)} a {n(a6['p_max_permutaciones'], 2)})."))
    R(4, ("Faltan mayo y julio; hogares incompletos algunos meses", "Faltan mayo y julio; agosto sin ficha; hogares incompletos"),
         ("Cruce de ambos bloques por niño y mes: 306 evaluaciones de 116 NNA con diagnóstico en la ficha",
          f"Cruce por periodo + ID (la fecha de nacimiento difiere entre bloques en {s2b['ninos_fecha_distinta']} NNA): {s3['evaluaciones']} evaluaciones de {s3['ninos_eval']} NNA con diagnóstico en la ficha"))
    R(5, ("Fechas en 3 formatos", "Fechas inconsistentes"),
         ("Texto mezclado día/mes", f"Ficha: 3 formatos. Bloque 2: 1 fecha mal transcrita en el consolidado anterior y {s2b['ninos_fecha_distinta']} NNA con nacimiento distinto entre bloques"),
         ("Conversión a fecha, 0 pérdidas", "Conversión por formato (0 pérdidas); consolidado rehecho desde los 6 Excel; cruce por periodo + ID"),
         ("14 columnas constantes o duplicadas", "13 variables constantes o duplicadas"),
         ("Valores faltantes", "Edad en el Bloque 2"),
         ("4 NNA sin tiempo en el hogar; 1 sin severidad", "Junio y agosto: fórmula HOY(), no la edad a la fecha de evaluación"),
         ("KNN (k=5) y mediana", "Edad calculada desde la fecha de nacimiento al día 15 del mes"))
    R(6, ("con 20 variables de la ficha", f"con {a4['predictores_finales']} variables de la ficha"),
         ("AUC entre 0,41 y 0,57; la prueba de permutaciones no los distingue del azar",
          f"AUC entre {n(a6['auc_min_modelos'], 2)} y {n(a6['auc_max_modelos'], 2)}; la prueba de permutaciones no los distingue del azar (p ≥ {n(a6['p_min_permutaciones'], 2)})"),
         ("La edad sola rinde igual o mejor que las 20 variables",
          f"La edad sola (AUC {n(a6['auc_solo_edad_estricta'], 2)}) rinde mejor que todas las variables ({n(a6['auc_todas_estricta'], 2)})"))
    cambiar_imagen(S[5], "Picture 5", FIG / "anexo6_viabilidad_auc.png")
    R(7, ("4 con KNN, 1 con mediana", "4 con KNN (tabla del Anexo; ajustado en train)"),
         ("Bloque 2: diagnósticos e IMC", "Bloque 2: consolidado desde 6 Excel"),
         ("379 filas; IMC recalculado (6 vacíos)", f"{s2a['evaluaciones']} filas; 19 hojas verificadas; IMC recalculado ({s2a['imc_recalculado']})"),
         ("306 filas, 116 NNA", f"{s3['evaluaciones']} filas, {s3['ninos_eval']} NNA (periodo + ID)"),
         ("concordancia 76 % en exceso de peso y 83 % en talla baja", f"concordancia {n(s3['conc_exceso'])} % en exceso de peso y {n(s3['conc_talla'])} % en talla baja"))
    R(8, ("(37,9 % de las evaluaciones)", f"({n(s3['prev_estricta'])} % de las evaluaciones)"),
         ("(68,0 %)", f"({n(s3['prev_amplia'])} %)"),
         ("87 niños, 229 evaluaciones", f"{s4['ninos_train']} niños, {s4['filas_train']} evaluaciones"),
         ("29 niños, 77 evaluaciones", f"{s4['ninos_test']} niños, {s4['filas_test']} evaluaciones"),
         ("Muestra pequeña. El conjunto de prueba tiene niños algo más pequeños que el de entrenamiento (edad media −0,3 desv.).",
          f"Muestra pequeña: el AUC cambia según qué niños caen en entrenamiento ({RANGO_LOGIT} en 20 particiones)."))
    R(9, ("0,87–0,94", f"{n(s4['corr_min_edad_peso_talla'], 2)}–{n(s4['corr_max_edad_peso_talla'], 2)}"),
         ("+2,5 kg, −1,4 cm y +1,5 de IMC", f"+{n(s4['res_peso_1'])} kg, {n(s4['res_talla_1'])} cm y +{n(s4['res_imc_1'])} de IMC".replace("-", "−")),
         ("−1,5 kg, +0,9 cm y −0,9 de IMC", f"{n(s4['res_peso_0'])} kg, +{n(s4['res_talla_0'])} cm y {n(s4['res_imc_0'])} de IMC".replace("-", "−")),
         ("StandardScaler ajustado con entrenamiento (KNN y SVC usan distancias)", "Residuos y StandardScaler dentro de un Pipeline: se ajustan en cada pliegue (sin fuga)"),
         ("CP1 = 85,5 %", f"CP1 = {n(s4['pca_cp1'])} %"))
    cambiar_imagen(S[8], "Picture 5", FIG / "step4_residuos_por_etiqueta.png")
    cambiar_imagen(S[8], "Picture 6", FIG / "step4_curvas_edad.png")
    R(10, ("Candidatos: regresión logística (interpretable) y SVC. KNN queda como comparador. Diferencias de ~0,02 entre logística y SVC no son concluyentes.",
           f"Candidatos: regresión logística (interpretable; supera al SVC en {s6['logit_gana_semillas']} de 20 particiones) y SVC. KNN queda como comparador."),
          ("0,81–0,82", f"{n(s6['auc_svc_ca'], 2)} ({RANGO_SVC})*"),
          ("0,82", f"{n(s6['auc_logit_ca'], 2)} ({RANGO_LOGIT})*"),
          ("0,62–0,78", f"{n(s6['auc_knn_min'], 2)}–{n(s6['auc_knn_max'], 2)}"),
          ("AUC real 0,81 vs 0,52 al azar (p ≈ 0,01)", f"AUC real {n(s6['auc_perm_real'], 2)} vs {n(s6['auc_perm_media'], 2)} al azar (p = {n(s6['p_perm'], 2)})"),
          ("La edad sola: AUC 0,52", f"La edad sola: AUC {n(s6['auc_solo_edad'], 2)}"),
          ("(AUC 0,62–0,72)", f"(AUC {n(s6['auc_amplia_min'], 2)}–{n(s6['auc_amplia_max'], 2)})"))
    agregar_vineta(S[9], "TextBox 9", "* Entre paréntesis: rango según 20 semillas de partición")
    cambiar_imagen(S[9], "Picture 5", FIG / "step6_auc_modelos.png")
    cambiar_imagen(S[9], "Picture 6", FIG / "step6_semillas.png")
    R(11, ("68,6 % de 306 evaluaciones.", f"{n(s5['concordancia'])} % de {s3['evaluaciones']} evaluaciones."),
          ("54 casos", f"{s5['ficha_sin_med_con']} casos"), ("42 casos", f"{s5['ficha_con_med_sin']} casos"),
          ("62 de 116 NNA con al menos una discrepancia (13 en todas sus evaluaciones).",
           f"{s5['ninos_discrepancia']} de {s5['ninos_eval']} NNA con al menos una discrepancia ({s5['ninos_todas']} en todas sus evaluaciones)."),
          ("(AUC 0,44–0,48)", f"(AUC {n(s6['km_crudas_min'], 2)}–{n(s6['km_crudas_max'], 2)})"),
          ("AUC hasta 0,74", f"AUC hasta {n(s6['km_ajust_max'], 2)}"))
    R(12, ("116 NNA, 2–5 evaluaciones cada uno", f"{s3['ninos_eval']} NNA, 1–5 evaluaciones cada uno"),
          ("Solo 6 de los 8 meses de 2026", "Solo 5 meses de 2026 con ficha (sin mayo ni julio; agosto sin ficha)"),
          ("La variabilidad entre repeticiones no refleja la incertidumbre real", f"AUC sensible a la partición ({RANGO_LOGIT} según la semilla)"),
          ("Logística y SVC no son distinguibles entre sí", "Junio: fecha de evaluación asumida (edad con fórmula HOY())"),
          ("Prueba (29 niños)", f"Prueba ({s4['ninos_test']} niños)"))
    R(13, ("– Revisión de los documentos del curso y de los datos para plantear el caso",
           "– Revisión de los documentos del curso y de los datos; auditoría del consolidado contra los Excel originales"),
          ("– Generación de código de los notebooks (limpieza, transformación y modelos) y de las interpretaciones iniciales",
           "– Reconstrucción del proyecto: código de los notebooks (Steps 1–6) y redacción de las interpretaciones, verificadas con asserts"),
          ("• Los prompts utilizados están disponibles como anexo a solicitud del docente",
           "• Declaración en cada notebook y declaración general con los prompts principales en el README.md"))
    prs.save(ENT / "Avance_TA_E10.pptx")


# ---------------------------------------------------------------- Ficha técnica
def ficha():
    wb = openpyxl.load_workbook(PLANT / "FichaTecnica_E10_avance_original.xlsx")
    ws = wb["FICHA TÉCNICA"]
    def R(celda, *pares):
        v = ws[celda].value
        for a, b in pares:
            if a not in v: raise SystemExit(f"No se encontró en {celda}: {a!r}")
            v = v.replace(a, b)
        ws[celda].value = v
    R("G28", ("164 NNA) y 379 evaluaciones antropométricas de 2026.",
              f"{s3['ninos_ficha']} NNA) y {s2a['evaluaciones']} evaluaciones antropométricas de 2026, consolidadas desde 6 Excel mensuales."))
    R("G32", ("concuerdan solo en 68,6 % de los casos", f"concuerdan solo en {n(s5['concordancia'])} % de las evaluaciones"))
    R("G40", ("(306 evaluaciones, 116 NNA)", f"({EVAL})"),
             ("14 NNA sin ningún diagnóstico, y 3 evaluaciones", f"evaluaciones de meses sin ficha (agosto), y 3 evaluaciones"))
    R("G44", ("Muestra pequeña (116 NNA, 2–5 evaluaciones cada uno); solo 6 de 8 meses de 2026 (faltan mayo y julio);",
              f"Muestra pequeña ({s3['ninos_eval']} NNA, 1–5 evaluaciones cada uno); solo 5 meses de 2026 con ficha (faltan mayo y julio; agosto sin ficha); AUC sensible a la partición ({RANGO_LOGIT});"),
             ("datos anonimizados.", "fecha de evaluación asumida al día 15 del mes (en junio la edad registrada es una fórmula); datos anonimizados."))
    R("G52", ("(62 de 116 con al menos una)", f"({s5['ninos_discrepancia']} de {s5['ninos_eval']} con al menos una)"),
             ("completar los meses faltantes;", "completar los meses faltantes; registrar la fecha de evaluación y la edad como valor (no como fórmula);"))
    R("G56", ("AUC ≈ 0,82 con regresión logística y SVC frente a 0,50 de la línea base. Concordancia ficha vs. medición OMS: 68,6 %.",
              f"AUC {n(s6['auc_logit_ca'], 2)} con regresión logística ({RANGO_LOGIT} según la partición) y {n(s6['auc_svc_ca'], 2)} con SVC, frente a 0,50 de la línea base. Concordancia ficha vs. medición OMS: {n(s5['concordancia'])} %."))
    wb.save(ENT / "FichaTecnica_E10.xlsx")


# ---------------------------------------------------------------- Gantt
def gantt():
    wb = openpyxl.load_workbook(PLANT / "Gantt_E10_avance_original.xlsx")
    ws = wb.active
    nombres = {5: "Steps 1–6 de KDD (avance)", 6: "Step 1–2: Dominio, consolidación y entendimiento de los datos", 7: "Step 3: Limpieza y preprocesamiento",
               8: "Step 4: Reducción y proyección", 9: "Step 5–6: Tarea de minería y selección de algoritmos (preliminar)",
               12: "Steps 7–9 de KDD (final)", 13: "Step 7: Ajuste de hiperparámetros (GridSearchCV agrupado)",
               14: "Step 7: Cuarto modelo (árbol/bosque) y comparación", 15: "Step 8: Evaluación final en el conjunto de prueba",
               16: "Step 9: Auditoría con la nutricionista (validar discrepancias)", 17: "Step 7: Clustering de perfiles (K-Prototypes / Gower)",
               18: "Step 8: Interpretación, conclusiones y líneas futuras", 19: "Step 9: Informe final, ficha técnica y notebooks finales"}
    for fila, texto in nombres.items(): ws.cell(fila, 1).value = texto
    ws.insert_rows(11)
    verde = ws.cell(10, 7).fill.fgColor.rgb                     # color de "hecho" (barra de la fila 10)
    for col in range(1, ws.max_column + 1):
        ws.cell(11, col)._style = copy.copy(ws.cell(10, col)._style)
        if col >= 5 and ws.cell(10, col).fill.fgColor.rgb == verde:   # quitar la barra copiada; dejar el fondo de la fila 12
            ws.cell(11, col).fill = copy.copy(ws.cell(12, col).fill)
    ws.cell(11, 1).value, ws.cell(11, 2).value, ws.cell(11, 3).value, ws.cell(11, 4).value = \
        "Reconstrucción y verificación del pipeline (Steps 1–6)", "29/10/2026", "29/10/2026", "Hecho"
    col_29 = next(c for c in range(5, ws.max_column + 1) if ws.cell(4, c).value == "29/10")
    ws.cell(11, col_29).fill = copy.copy(ws.cell(10, 7).fill)  # verde = hecho
    ws["A2"].value = ws["A2"].value.replace("Avance al 28/10/2026.", "Avance al 29/10/2026.")
    wb.save(ENT / "Gantt_E10.xlsx")


if __name__ == "__main__":
    pptx(); ficha(); gantt()
    print("OK – entregables generados en", ENT)
